#!/usr/bin/env python3
"""Build leakage-safe counter-site rows with real labels and raw spatial features."""

import csv
from collections import defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import random
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print("Install data dependencies first: python3 -m pip install -r requirements-data.txt", file=sys.stderr)
    raise

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"
CONFIG = json.loads((ROOT / "data" / "training-config.json").read_text())
R_EARTH = 6_371_000


def haversine(lon1, lat1, lon2, lat2):
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R_EARTH * math.asin(math.sqrt(a))


def segment_length(a, b):
    return haversine(a[0], a[1], b[0], b[1])


def point_segment_distance(point, a, b):
    lon, lat = point
    scale_x = 111_320 * math.cos(math.radians(lat))
    scale_y = 110_540
    ax, ay = (a[0] - lon) * scale_x, (a[1] - lat) * scale_y
    bx, by = (b[0] - lon) * scale_x, (b[1] - lat) * scale_y
    dx, dy = bx - ax, by - ay
    if dx == 0 and dy == 0:
        return math.hypot(ax, ay)
    t = max(0, min(1, -(ax * dx + ay * dy) / (dx * dx + dy * dy)))
    return math.hypot(ax + t * dx, ay + t * dy)


def lines(geometry):
    if geometry["type"] == "LineString":
        return [geometry["coordinates"]]
    if geometry["type"] == "MultiLineString":
        return geometry["coordinates"]
    return []


def point_in_ring(point, ring):
    x, y = point
    inside = False
    j = len(ring) - 1
    for i, (xi, yi, *_) in enumerate(ring):
        xj, yj, *_ = ring[j]
        if (yi > y) != (yj > y):
            crossing = (xj - xi) * (y - yi) / (yj - yi) + xi
            if x < crossing:
                inside = not inside
        j = i
    return inside


def point_in_geometry(point, geometry):
    polygons = [geometry["coordinates"]] if geometry["type"] == "Polygon" else geometry["coordinates"]
    return any(point_in_ring(point, polygon[0]) and not any(point_in_ring(point, hole) for hole in polygon[1:]) for polygon in polygons)


def census_profiles():
    workbook = load_workbook(RAW / "neighbourhood-profiles-2021.xlsx", read_only=True, data_only=True)
    sheet = workbook["hd2021_census_profile"]
    rows = list(sheet.iter_rows(values_only=True))
    numbers = rows[1][1:]
    wanted = {
        "population": "Total - Age groups of the population - 25% sample data",
        "low_income_pct": "Prevalence of low income based on the Low-income measure, after tax (LIM-AT) (%)",
    }
    values = {}
    for key, label in wanted.items():
        match = next((row for row in rows if str(row[0]).strip() == label), None)
        if match is None:
            raise ValueError(f"Missing census indicator: {label}")
        values[key] = match[1:]
    return {
        str(number): {key: values[key][index] for key in wanted}
        for index, number in enumerate(numbers)
    }


def build_spatial_features(location_features=None):
    required = [
        "counter-intervals-2024-2025.csv", "bike-counter-locations.geojson", "flows.json",
        "ksi-collisions.geojson", "cycling-network.geojson",
        "neighbourhood-boundaries.geojson", "neighbourhood-profiles-2021.xlsx",
    ]
    paths = {
        name: (ROOT / "public" / "data" / name if name == "flows.json" else RAW / name)
        for name in required
    }
    missing = [name for name, path in paths.items() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing required inputs: {', '.join(missing)}")

    counter_locations = json.loads(paths["bike-counter-locations.geojson"].read_text())["features"]
    if location_features is not None:
        counter_locations = location_features
    location_by_id = {str(f["properties"]["location_dir_id"]): f for f in counter_locations}
    bike_share = json.loads(paths["flows.json"].read_text())["stations"]
    network = json.loads(paths["cycling-network.geojson"].read_text())["features"]
    collisions_raw = json.loads(paths["ksi-collisions.geojson"].read_text())["features"]
    cyclist_collisions = {}
    for feature in collisions_raw:
        props = feature["properties"]
        if not "2019-06-01" <= str(props.get("accdate", ""))[:10] < "2024-06-01":
            continue
        if str(props.get("cyclist", "")).lower() != "true" and str(props.get("road_user", "")).lower() != "cyclist":
            continue
        longitude, latitude = props.get("longitude"), props.get("latitude")
        if isinstance(longitude, (int, float)) and isinstance(latitude, (int, float)):
            cyclist_collisions[str(props.get("collision_id"))] = [longitude, latitude]

    boundaries = json.loads(paths["neighbourhood-boundaries.geojson"].read_text())["features"]
    profiles = census_profiles()
    radius = CONFIG["spatialRadiiMetres"]
    static = {}

    for location_id, feature in location_by_id.items():
        point = feature["geometry"]["coordinates"][:2]
        nearby_stations = [s for s in bike_share if haversine(point[0], point[1], s["lon"], s["lat"]) <= radius["bikeShare"]]
        collision_count = sum(haversine(point[0], point[1], c[0], c[1]) <= radius["cyclistKsi"] for c in cyclist_collisions.values())
        network_metres = 0.0
        for segment in network:
            for line in lines(segment["geometry"]):
                for a, b in zip(line, line[1:]):
                    if point_segment_distance(point, a, b) <= radius["cyclingNetwork"]:
                        network_metres += segment_length(a, b)
        area = next((f for f in boundaries if point_in_geometry(point, f["geometry"])), None)
        area_code = str(int(area["properties"]["AREA_SHORT_CODE"])) if area else None
        profile = profiles.get(area_code, {})
        static[location_id] = {
            "counter_site": feature["properties"]["location_name"],
            "longitude": point[0],
            "latitude": point[1],
            "bike_share_stations_1km": len(nearby_stations),
            "bike_share_trip_ends_1km": sum(s["trip_ends_total"] for s in nearby_stations),
            "cyclist_ksi_collisions_1km": collision_count,
            "cycling_network_km_1km": round(network_metres / 1000, 4),
            "neighbourhood_id": area_code,
            "neighbourhood_population": profile.get("population"),
            "neighbourhood_low_income_pct": profile.get("low_income_pct"),
            "neighbourhood_improvement_area": int(area is not None and area["properties"].get("CLASSIFICATION_CODE") == "NIA"),
        }

    return static, location_by_id, paths


def main():
    static, location_by_id, paths = build_spatial_features()
    start = date.fromisoformat(CONFIG["labelPeriod"]["start"])
    end = date.fromisoformat(CONFIG["labelPeriod"]["endExclusive"])
    rows = []
    daily = defaultdict(lambda: [0.0, 0.0])
    seen_bins = set()
    with paths["counter-intervals-2024-2025.csv"].open(newline="") as handle:
        for record in csv.DictReader(handle):
            observed_date = date.fromisoformat(record["datetime_bin"][:10])
            if not start <= observed_date < end:
                continue
            location_id = str(record["location_dir_id"])
            if location_id not in location_by_id or not record["bin_volume"]:
                continue
            volume = float(record["bin_volume"])
            key = (location_id, record["datetime_bin"])
            if key in seen_bins or volume < 0:
                continue
            seen_bins.add(key)
            hh, mm, ss = map(int, location_by_id[location_id]["properties"]["bin_size"].split(":"))
            hours = hh + mm / 60 + ss / 3600
            if not 0 < hours <= 1:
                raise ValueError("Unsupported counter bin size")
            daily[(location_id, observed_date)][0] += volume
            daily[(location_id, observed_date)][1] += hours
    for (location_id, observed_date), (volume, hours) in sorted(daily.items()):
            if hours < 20 or hours > 24 or volume > 10000:
                continue
            features = static.get(location_id)
            if not features or features["neighbourhood_population"] is None:
                continue
            group_id = hashlib.sha256(features["counter_site"].encode()).hexdigest()[:12]
            rows.append({
                "observation_date": observed_date.isoformat(),
                "group_id": group_id,
                "location_direction_id": location_id,
                "direction": location_by_id[location_id]["properties"]["direction"],
                "observed_hours": hours,
                "bicycles_per_observed_hour": round(volume / hours, 6),
                **features,
                "day_of_week": observed_date.weekday(),
            })

    groups = sorted({row["group_id"] for row in rows})
    shuffled = groups[:]
    random.Random(CONFIG["splitSeed"]).shuffle(shuffled)
    test_count = max(1, round(len(groups) * CONFIG["testFraction"]))
    test_groups = sorted(shuffled[:test_count])
    train_groups = sorted(set(groups) - set(test_groups))
    if set(train_groups) & set(test_groups):
        raise AssertionError("Spatial groups overlap")

    PROCESSED.mkdir(parents=True, exist_ok=True)
    columns = list(rows[0]) if rows else []
    with (PROCESSED / "counter-training.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)

    manifest = {
        "schemaVersion": "velocity.training-data.v1",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "featureVersion": CONFIG["featureVersion"],
        "target": CONFIG["target"],
        "rowCount": len(rows),
        "trainingDataSha256": hashlib.sha256((PROCESSED / "counter-training.csv").read_bytes()).hexdigest(),
        "configurationSha256": hashlib.sha256((ROOT / "data/training-config.json").read_bytes()).hexdigest(),
        "inputFileHashes": {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in paths.values()},
        "grouping": CONFIG["grouping"],
        "trainGroupIds": train_groups,
        "testGroupIds": test_groups,
        "normalization": CONFIG["normalization"],
        "observationStart": CONFIG["labelPeriod"]["start"],
        "observationEnd": (end.fromordinal(end.toordinal() - 1)).isoformat(),
        "sourceIds": [
            "permanent-bike-counter-intervals", "permanent-bike-counter-locations",
            "bike-share-trips-2024", "bike-share-stations", "cycling-network",
            "ksi-collisions", "neighbourhood-profiles-2021", "neighbourhood-boundaries-158",
        ],
        "warnings": [
            "Labels use summed valid interval counts divided by observed hours; days below 20 hours are excluded. Zero counts retained; obstruction can depress counts. Directional daily mean, not peak-hour demand.",
            "Collision features use deduplicated cyclist-involved KSI events from 2019-06-01 through 2024-05-31.",
            "2026 station and cycling-network snapshots accompany 2024 observations; spatial backtesting is not historical forecasting validation.",
            "Bike Share station activity reflects June 2024 station coverage and users, not all cycling demand.",
            "Static spatial features are repeated across daily observations and must be grouped by counter site during validation."
        ],
    }
    (PROCESSED / "training-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"Wrote {len(rows)} rows across {len(train_groups)} train and {len(test_groups)} test counter-site groups")


if __name__ == "__main__":
    main()
